"""Test lane AI tanpa jaringan: pipeline extract -> plan -> generate -> verify
-> strip, dan Penilaian AI (DESIGN.md §6).

Client LLM dipalsukan lewat `monkeypatch` pada `model_factory.build_client`,
jadi tidak ada panggilan HTTP sama sekali. Dokumen uji dibuat sebagai PPTX asli
(lalu di-ekstraksi `python-pptx` sungguhan) supaya fast path ikut teruji.
"""

from __future__ import annotations

import asyncio
import logging
from dataclasses import dataclass, field
from pathlib import Path

import pytest
from conftest import FAKE_PDF, make_session_row, make_tugas, seed_questions
from pydantic import BaseModel

from api.agents import model_factory, pipeline, tasks
from api.agents.assess import AnswerPair, assess_document
from api.agents.extract import MAX_CONTEXT_CHARS, extract_document, truncate_for_context
from api.agents.prompts import INJECTION_GUARD, _SYSTEM_ID
from api.agents.schemas import (
    DraftQuestion,
    PenilaianAIDraft,
    QuestionPlan,
    question_set_schema,
)
from api.core.enums import QuestionType, SessionStatus
from api.core.models.penilaian_ai import PenilaianAI
from api.core.models.pertanyaan import Jawaban, Pertanyaan
from api.core.models.session import SessionRow
from api.core.models.unggahan import Unggahan

MIME_PDF = "application/pdf"
MIME_PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

BRIEF = "Jelaskan analisis ekonomi digital beserta dampaknya bagi UMKM lokal."

# Paragraf materi — dipakai untuk membangun PPTX uji dan sebagai kutipan
# evidence yang harus lolos verifikasi substring.
DOC_PARAGRAPHS: tuple[str, ...] = (
    "Analisis ekonomi digital menjelaskan bagaimana aktivitas ekonomi berpindah ke platform daring.",
    "Data diambil dari survei 120 UMKM di Bandung pada bulan Maret 2024.",
    "Temuan utama: 64 persen UMKM menaikkan omzet lewat kanal marketplace.",
    "Hambatan terbesar adalah biaya logistik dan kurangnya tenaga digital.",
    "Rekomendasi: pelatihan pembukuan sederhana dan pendampingan unggah katalog.",
    "Dampak lanjutan mencakup lahirnya pekerjaan baru di bidang pengelolaan katalog.",
)


# --------------------------------------------------------------------------- #
# Fake client LLM
# --------------------------------------------------------------------------- #
@dataclass
class Call:
    """Satu panggilan yang tercatat, buat assert isi prompt di test."""

    schema_name: str
    prompt: str
    inline_parts: tuple = ()


@dataclass
class FakeClient:
    """Client terstruktur palsu: mengembalikan respons berurutan dari antrean."""

    responses: list
    supports_inline_data: bool = False
    calls: list[Call] = field(default_factory=list)

    async def complete(self, *, prompt, schema, instruction=None, inline_parts=None):
        self.calls.append(
            Call(
                schema_name=schema.__name__,
                prompt=prompt,
                inline_parts=tuple(inline_parts or ()),
            )
        )
        if not self.responses:
            raise AssertionError(f"fake client kehabisan respons (schema={schema.__name__})")
        nxt = self.responses.pop(0)
        if isinstance(nxt, Exception):
            raise nxt
        if isinstance(nxt, BaseModel):
            return nxt
        return schema.model_validate(nxt)


def plan_dict(
    *,
    ideasi_angle="asal ide dari survei lapangan di Bandung",
    isi_angle="angka hasil survei 120 UMKM bulan Maret 2024",
    harapan_angle="dampak lanjutan pelatihan pembukuan ke UMKM",
):
    """Plan valid: satu template per tipe, tiga angle berbeda."""
    return {
        "templates": [
            {
                "type": QuestionType.IDEASI.value,
                "angle": ideasi_angle,
                "alasan": "Asal ide tercermin langsung di dalam materi.",
            },
            {
                "type": QuestionType.ISI.value,
                "angle": isi_angle,
                "alasan": "Angka di materi membuktikan pengerjaan sendiri.",
            },
            {
                "type": QuestionType.HARAPAN.value,
                "angle": harapan_angle,
                "alasan": "Materi menyebut langkah lanjutan yang konkret.",
            },
        ]
    }


def questions_dict(evidence_by_type: dict[str, str]):
    """QuestionSet dengan tiga pertanyaan; evidence ditentukan test."""
    texts = {
        QuestionType.IDEASI.value: "Dari mana ide utama karya ini berasal?",
        QuestionType.ISI.value: "Data apa yang kamu kumpulkan sendiri dalam karya ini?",
        QuestionType.HARAPAN.value: "Apa langkah berikutnya yang kamu rencanakan?",
    }
    return {
        "questions": [
            {"type": tipe, "text": texts[tipe], "evidence": evidence_by_type[tipe]}
            for tipe in (
                QuestionType.IDEASI.value,
                QuestionType.ISI.value,
                QuestionType.HARAPAN.value,
            )
        ]
    }


VALID_EVIDENCE = {
    tipe.value: DOC_PARAGRAPHS[index] for index, tipe in enumerate(QuestionType)
}


# --------------------------------------------------------------------------- #
# Fixture
# --------------------------------------------------------------------------- #
def _write_pptx(path: Path) -> None:
    """Buat PPTX asli dengan teks yang cukup panjang (bukan materi tipis)."""
    from pptx import Presentation

    presentation = Presentation()
    for index in range(0, len(DOC_PARAGRAPHS), 3):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"Bagian {index // 3 + 1} Analisis Ekonomi Digital"
        body = slide.placeholders[1]
        body.text_frame.text = DOC_PARAGRAPHS[index]
        for extra in DOC_PARAGRAPHS[index + 1 : index + 3]:
            body.text_frame.add_paragraph().text = extra
    presentation.save(str(path))


@pytest.fixture()
def pptx_path(tmp_path: Path) -> Path:
    path = tmp_path / "karya.pptx"
    _write_pptx(path)
    return path


def _attach_unggahan(db, row: SessionRow, path: Path, mime: str, filename: str) -> SessionRow:
    with db() as session:
        session.add(
            Unggahan(
                session_id=row.id,
                filename=filename,
                mime=mime,
                bytes=path.stat().st_size,
                storage_path=str(path),
            )
        )
        session.commit()
    return row


@pytest.fixture()
def sesi_dengan_pptx(db, guru, murid, pptx_path):
    """Sesi di status `generating` + Unggahan PPTX di disk."""
    with db() as session:
        tugas = make_tugas(guru, session, brief=BRIEF)
        row = make_session_row(session, tugas, murid, status=SessionStatus.GENERATING)
        session.commit()
    return _attach_unggahan(db, row, pptx_path, MIME_PPTX, "karya.pptx")


@pytest.fixture()
def fake_model(monkeypatch):
    """Pasang client palsu ke pabrik model."""

    def _install(client: FakeClient) -> FakeClient:
        monkeypatch.setattr(model_factory, "build_client", lambda: client)
        return client

    return _install


def _questions(db, session_id: str) -> list[Pertanyaan]:
    with db() as session:
        return list(
            session.query(Pertanyaan)
            .filter(Pertanyaan.session_id == session_id)
            .order_by(Pertanyaan.seq)
            .all()
        )


def _status(db, session_id: str) -> SessionStatus:
    with db() as session:
        return session.get(SessionRow, session_id).status


def _seed_submitted(db, row: SessionRow, *, with_answers: bool = True) -> list[str]:
    """Isi Pertanyaan + Jawaban lalu paksa status `submitted`."""
    with db() as session:
        ids = seed_questions(session, row.id, with_evidence=True)
        if with_answers:
            for index, pertanyaan_id in enumerate(ids):
                session.add(
                    Jawaban(
                        pertanyaan_id=pertanyaan_id,
                        text=f"jawaban murid nomor {index + 1} yang cukup panjang untuk dinilai",
                    )
                )
        session.get(SessionRow, row.id).status = SessionStatus.SUBMITTED
        session.commit()
    return ids


# --------------------------------------------------------------------------- #
# Pipeline murni
# --------------------------------------------------------------------------- #
def test_plan_requires_distinct_angles():
    templates = [
        (QuestionType.IDEASI, "asal ide dari survei", "alasan"),
        (QuestionType.ISI, "Asal Ide dari Survei", "alasan"),
        (QuestionType.HARAPAN, "dampak lanjutan", "alasan"),
    ]
    reason = pipeline.validate_plan(templates)
    assert reason is not None
    assert "harus berbeda" in reason


def test_plan_valid_when_angles_differ():
    templates = [
        (QuestionType.IDEASI, "asal ide dari survei lapangan", "alasan"),
        (QuestionType.ISI, "angka hasil survei 120 UMKM", "alasan"),
        (QuestionType.HARAPAN, "dampak lanjutan pelatihan pembukuan", "alasan"),
    ]
    assert pipeline.validate_plan(templates) is None


def test_plan_rejects_duplicate_type():
    templates = [
        (QuestionType.IDEASI, "asal ide dari survei lapangan", "alasan"),
        (QuestionType.IDEASI, "angka hasil survei 120 UMKM", "alasan"),
        (QuestionType.HARAPAN, "dampak lanjutan pelatihan pembukuan", "alasan"),
    ]
    reason = pipeline.validate_plan(templates)
    assert reason is not None
    assert "jumlah template per tipe harus 1" in reason


def test_plan_supports_more_than_one_question_per_type():
    templates = [
        (QuestionType.IDEASI, "asal ide dari survei lapangan", "alasan"),
        (QuestionType.IDEASI, "referensi yang dipakai saat menyusun ide", "alasan"),
        (QuestionType.ISI, "angka hasil survei 120 UMKM", "alasan"),
        (QuestionType.ISI, "pilihan layout halaman presentasi", "alasan"),
        (QuestionType.HARAPAN, "dampak lanjutan pelatihan pembukuan", "alasan"),
        (QuestionType.HARAPAN, "targetjai tahun depan", "alasan"),
    ]
    assert pipeline.validate_plan(templates, 2) is None
    assert pipeline.validate_plan(templates, 1) is not None


def test_evidence_verification_normalizes_case_bullets_and_line_numbers():
    document = "\n".join(
        [
            "## Slide 1",
            "1. Analisis ekonomi digital menjelaskan aktivitas ekonomi daring.",
            "- Data diambil dari survei 120 UMKM di Bandung.",
        ]
    )
    assert pipeline.verify_evidence("- Data diambil dari Survei 120  UMKM di bandung", document)
    assert not pipeline.verify_evidence("Data diambil dari survei 500 UMKM", document)
    assert not pipeline.verify_evidence("pendulum gravitasi newton", document)
    assert not pipeline.verify_evidence("pendulum gravitasi newton", "")


def test_truncate_for_context_caps_at_limit():
    text = "x" * (MAX_CONTEXT_CHARS + 500)
    capped, truncated = truncate_for_context(text)
    assert truncated is True
    assert len(capped) < len(text)
    assert "MATERI DIPOTONG" in capped

    same, not_truncated = truncate_for_context("pendek")
    assert (same, not_truncated) == ("pendek", False)


def test_system_prompt_contains_injection_guard_verbatim():
    assert INJECTION_GUARD in _SYSTEM_ID
    assert "TIDAK trustworthy" in _SYSTEM_ID


def test_question_set_schema_locks_count():
    tiga = question_set_schema(3)
    with pytest.raises(Exception):
        tiga.model_validate(
            {"questions": [{"type": "ideasi", "text": "a" * 20, "evidence": "b" * 20}]}
        )

    satu = question_set_schema(1)
    parsed = satu.model_validate(
        {"questions": [{"type": "isi", "text": "a" * 20, "evidence": "b" * 20}]}
    )
    assert parsed.questions[0].type == QuestionType.ISI


def test_draft_question_schema_is_strict_about_evidence_length():
    with pytest.raises(Exception):
        DraftQuestion(type=QuestionType.ISI, text="a" * 30, evidence="pendek")
    with pytest.raises(Exception):
        PenilaianAIDraft.model_validate(
            {
                "skor_ideasi": "bukan angka",
                "skor_isi": 50,
                "skor_harapan": 50,
                "ringkasan": "r" * 30,
            }
        )


def test_clamp_score_holds_0_to_100():
    from api.agents.assess import clamp_score

    clamped: list[str] = []
    assert clamp_score(80, "skor_ideasi", clamped=clamped) == 80
    assert clamp_score(150, "skor_ideasi", clamped=clamped) == 100
    assert clamp_score(-20, "skor_isi", clamped=clamped) == 0
    assert clamped == ["skor_ideasi", "skor_isi"]
    assert clamp_score("70", "skor_harapan") == 70
    with pytest.raises(Exception):
        clamp_score("baik sekali", "skor_isi")


def test_plan_schema_requires_three_templates():
    with pytest.raises(Exception):
        QuestionPlan.model_validate(plan_dict()["templates"][:2])


# --------------------------------------------------------------------------- #
# extract
# --------------------------------------------------------------------------- #
def test_extract_pptx_fast_path_keeps_document_text(pptx_path):
    document = extract_document(
        Unggahan(
            session_id="x" * 32,
            filename="karya.pptx",
            mime=MIME_PPTX,
            bytes=pptx_path.stat().st_size,
            storage_path=str(pptx_path),
        )
    )

    assert document.fast_path == "python-pptx"
    assert document.error is None
    assert document.is_thin is False
    haystack = pipeline.normalize_for_match(document.markdown)
    for paragraph in DOC_PARAGRAPHS:
        assert pipeline.normalize_for_match(paragraph) in haystack


def test_extract_broken_pdf_falls_back_to_thin_text(tmp_path):
    path = tmp_path / "rusak.pdf"
    path.write_bytes(FAKE_PDF)

    document = extract_document(
        Unggahan(
            session_id="y" * 32,
            filename="rusak.pdf",
            mime=MIME_PDF,
            bytes=path.stat().st_size,
            storage_path=str(path),
        )
    )

    assert document.fast_path == "markitdown"
    assert document.is_thin is True
    assert document.inline_data() is not None


# --------------------------------------------------------------------------- #
# generate_questions (entry point)
# --------------------------------------------------------------------------- #
def test_generate_questions_happy_path_writes_evidence_and_transitions(
    db, sesi_dengan_pptx, fake_model
):
    client = fake_model(FakeClient([plan_dict(), questions_dict(VALID_EVIDENCE)]))

    tasks.generate_questions(sesi_dengan_pptx.id)

    rows = _questions(db, sesi_dengan_pptx.id)
    assert [row.type for row in rows] == list(QuestionType)
    assert [row.seq for row in rows] == [0, 1, 2]
    assert all(row.evidence in DOC_PARAGRAPHS for row in rows)
    assert all(row.text.strip() for row in rows)
    assert _status(db, sesi_dengan_pptx.id) == SessionStatus.ANSWERING
    assert [call.schema_name for call in client.calls] == ["QuestionPlan", "QuestionSet"]
    assert BRIEF in client.calls[0].prompt
    assert DOC_PARAGRAPHS[0][:40] in client.calls[0].prompt


def test_generate_questions_repairs_fabricated_evidence(db, sesi_dengan_pptx, fake_model):
    fabricated = dict(VALID_EVIDENCE)
    fabricated[QuestionType.ISI.value] = "Survei dilakukan terhadap 900 UMKM di Jakarta."
    repaired = {
        "questions": [
            {
                "type": QuestionType.ISI.value,
                "text": "Data apa yang kamu kumpulkan sendiri dalam karya ini?",
                "evidence": DOC_PARAGRAPHS[1],
            }
        ]
    }
    client = fake_model(FakeClient([plan_dict(), questions_dict(fabricated), repaired]))

    tasks.generate_questions(sesi_dengan_pptx.id)

    isi = next(
        row for row in _questions(db, sesi_dengan_pptx.id) if row.type == QuestionType.ISI
    )
    assert isi.evidence == DOC_PARAGRAPHS[1]
    assert [call.schema_name for call in client.calls] == [
        "QuestionPlan",
        "QuestionSet",
        "QuestionSet",
    ]
    assert "tidak ditemukan di dalam dokumen" in client.calls[2].prompt
    assert _status(db, sesi_dengan_pptx.id) == SessionStatus.ANSWERING


def test_generate_questions_replans_when_angles_collide(db, sesi_dengan_pptx, fake_model):
    bentrok = plan_dict(isi_angle="asal ide dari survei lapangan di Bandung")
    client = fake_model(FakeClient([bentrok, plan_dict(), questions_dict(VALID_EVIDENCE)]))

    tasks.generate_questions(sesi_dengan_pptx.id)

    assert [call.schema_name for call in client.calls] == [
        "QuestionPlan",
        "QuestionPlan",
        "QuestionSet",
    ]
    assert "harus berbeda" in client.calls[1].prompt
    assert len(_questions(db, sesi_dengan_pptx.id)) == 3
    assert _status(db, sesi_dengan_pptx.id) == SessionStatus.ANSWERING


def test_generate_questions_keeps_best_effort_after_failed_repairs(
    db, sesi_dengan_pptx, fake_model
):
    fabricated = dict(VALID_EVIDENCE)
    fabricated[QuestionType.HARAPAN.value] = "kutipan pertama yang mengarang data baru"
    gagal = {
        "questions": [
            {
                "type": QuestionType.HARAPAN.value,
                "text": "Apa langkah berikutnya yang kamu rencanakan?",
                "evidence": "kutipan kedua yang juga mengarang data baru",
            }
        ]
    }
    client = fake_model(
        FakeClient([plan_dict(), questions_dict(fabricated), gagal, dict(gagal)])
    )

    tasks.generate_questions(sesi_dengan_pptx.id)

    harapan = next(
        row for row in _questions(db, sesi_dengan_pptx.id) if row.type == QuestionType.HARAPAN
    )
    # Best effort: pertanyaan tetap tersimpan supaya Guru bisa mengauditnya.
    assert harapan.evidence
    assert len(client.calls) == 4
    assert _status(db, sesi_dengan_pptx.id) == SessionStatus.ANSWERING


def test_generate_questions_scales_with_questions_per_type(
    db, pptx_path, guru, murid, fake_model
):
    """`questions_per_type=2` -> enam Pertanyaan (harus sesuai `expected_question_count`)."""
    with db() as session:
        tugas = make_tugas(guru, session, brief=BRIEF, questions_per_type=2)
        row = make_session_row(session, tugas, murid, status=SessionStatus.GENERATING)
        session.commit()
    _attach_unggahan(db, row, pptx_path, MIME_PPTX, "karya.pptx")

    plan_enam = {
        "templates": [
            {
                "type": tipe.value,
                "angle": f"{tipe.value} sudut pandang nomor {slot + 1}",
                "alasan": f"Sudut {tipe.value} nomor {slot + 1} didukung materi.",
            }
            for slot in range(2)
            for tipe in QuestionType
        ]
    }
    soal = {
        "questions": [
            {
                "type": tipe.value,
                "text": f"Pertanyaan {tipe.value} nomor {slot + 1} yang cukup panjang",
                "evidence": DOC_PARAGRAPHS[(slot * 3) % len(DOC_PARAGRAPHS)],
            }
            for slot in range(2)
            for tipe in QuestionType
        ]
    }
    fake_model(FakeClient([plan_enam, soal]))

    tasks.generate_questions(row.id)

    rows = _questions(db, row.id)
    assert len(rows) == 6
    assert [row.seq for row in rows] == [0, 1, 2, 3, 4, 5]
    assert _status(db, row.id) == SessionStatus.ANSWERING


def test_generate_questions_is_idempotent(db, sesi_dengan_pptx, fake_model):
    client = fake_model(FakeClient([plan_dict(), questions_dict(VALID_EVIDENCE)]))

    tasks.generate_questions(sesi_dengan_pptx.id)
    first = [row.id for row in _questions(db, sesi_dengan_pptx.id)]

    tasks.generate_questions(sesi_dengan_pptx.id)

    assert len(client.calls) == 2, "panggilan kedua tidak boleh memanggil model"
    assert [row.id for row in _questions(db, sesi_dengan_pptx.id)] == first


def test_generate_questions_swallows_exception_and_keeps_session_alive(
    db, sesi_dengan_pptx, fake_model, caplog
):
    fake_model(FakeClient([RuntimeError("provider tumbang")]))

    with caplog.at_level(logging.ERROR):
        tasks.generate_questions(sesi_dengan_pptx.id)  # tidak melempar

    assert _questions(db, sesi_dengan_pptx.id) == []
    assert _status(db, sesi_dengan_pptx.id) == SessionStatus.GENERATING
    assert any("generate_questions gagal" in record.message for record in caplog.records)


def test_generate_questions_swallows_plan_error_after_retry(
    db, sesi_dengan_pptx, fake_model, caplog
):
    bentrok = plan_dict(isi_angle="asal ide dari survei lapangan di Bandung")
    fake_model(FakeClient([bentrok, bentrok]))

    with caplog.at_level(logging.ERROR):
        tasks.generate_questions(sesi_dengan_pptx.id)

    assert _questions(db, sesi_dengan_pptx.id) == []
    assert _status(db, sesi_dengan_pptx.id) == SessionStatus.GENERATING
    assert any("generate_questions gagal" in record.message for record in caplog.records)


def test_generate_questions_skipped_without_unggahan(db, guru, murid, fake_model):
    with db() as session:
        tugas = make_tugas(guru, session)
        row = make_session_row(session, tugas, murid, status=SessionStatus.GENERATING)
        session.commit()

    client = fake_model(FakeClient([]))
    tasks.generate_questions(row.id)

    assert client.calls == []
    assert _status(db, row.id) == SessionStatus.GENERATING


def test_generate_questions_skipped_when_session_terminal(db, sesi_dengan_pptx, fake_model):
    with db() as session:
        session.get(SessionRow, sesi_dengan_pptx.id).status = SessionStatus.AUTO_SUBMITTED
        session.commit()

    client = fake_model(FakeClient([]))
    tasks.generate_questions(sesi_dengan_pptx.id)

    assert client.calls == []
    assert _questions(db, sesi_dengan_pptx.id) == []


def test_thin_document_sends_inline_data_only_when_supported(
    db, tmp_path, guru, murid, fake_model
):
    """Fallback multimodal: PDF rusak -> teks tipis; binary dikirim sebagai part."""
    path = tmp_path / "rusak.pdf"
    path.write_bytes(FAKE_PDF)
    with db() as session:
        tugas = make_tugas(guru, session)
        row = make_session_row(session, tugas, murid, status=SessionStatus.GENERATING)
        session.commit()
    _attach_unggahan(db, row, path, MIME_PDF, "rusak.pdf")

    client = fake_model(
        FakeClient([plan_dict(), questions_dict(VALID_EVIDENCE)], supports_inline_data=True)
    )
    tasks.generate_questions(row.id)

    # Materi dibaca sebagai binary -> substring check tidak bisa dilakukan;
    # verifikasi dilewati dan evidence disimpan apa adanya untuk audit Guru.
    assert client.calls[0].inline_parts, "materi tipis harus dikirim sebagai inline_data"
    assert client.calls[0].inline_parts[0].mime_type == MIME_PDF
    assert len(client.calls) == 2, "tidak ada repair call saat verifikasi dilewati"
    assert len(_questions(db, row.id)) == 3


def test_thin_document_skips_inline_data_when_model_is_text_only(
    db, tmp_path, guru, murid, fake_model, caplog
):
    path = tmp_path / "rusak.pdf"
    path.write_bytes(FAKE_PDF)
    with db() as session:
        tugas = make_tugas(guru, session)
        row = make_session_row(session, tugas, murid, status=SessionStatus.GENERATING)
        session.commit()
    _attach_unggahan(db, row, path, MIME_PDF, "rusak.pdf")

    client = fake_model(FakeClient([plan_dict(), questions_dict(VALID_EVIDENCE)]))

    with caplog.at_level(logging.WARNING):
        tasks.generate_questions(row.id)

    assert all(call.inline_parts == () for call in client.calls)
    assert any("tidak mendukung inline data" in record.message for record in caplog.records)
    assert any("verifikasi evidence dilewati" in record.message for record in caplog.records)
    assert len(_questions(db, row.id)) == 3


# --------------------------------------------------------------------------- #
# Integrasi: upload -> BackgroundTasks -> ai_hooks -> tasks
# --------------------------------------------------------------------------- #
def test_upload_triggers_generate_through_background_tasks(
    client, login_as, murid, guru, db, fake_model
):
    """`POST /unggahan` menjadwalkan `generate_questions` lewat ai_hooks.

    TestClient menyelesaikan BackgroundTasks sebelum mengembalikan response,
    jadi tiga Pertanyaan harus sudah ada setelah POST selesai.
    """
    fake_model(FakeClient([plan_dict(), questions_dict(VALID_EVIDENCE)]))
    login_as(murid)
    with db() as session:
        tugas = make_tugas(guru, session, brief=BRIEF)
    session_id = client.post("/api/sessions", json={"tugas_id": tugas.id}).json()["id"]

    response = client.post(
        f"/api/sessions/{session_id}/unggahan",
        files={"file": ("karya.pdf", FAKE_PDF, MIME_PDF)},
    )
    assert response.status_code == 200

    rows = _questions(db, session_id)
    assert [row.type for row in rows] == list(QuestionType)

    # Sesi sudah `answering`; endpoint status Murid tidak pernah bawa skor.
    status = client.get(f"/api/sessions/{session_id}/status").json()
    assert status["status"] == "answering"
    assert "skor" not in str(status)

    # Payload Murid tidak pernah membawa evidence.
    questions = client.get(f"/api/sessions/{session_id}/questions").json()
    assert len(questions["questions"]) == 3
    assert "evidence" not in str(questions)


# --------------------------------------------------------------------------- #
# assess_answers (entry point)
# --------------------------------------------------------------------------- #
def test_assess_answers_recomputes_total_and_ignores_declared(
    db, sesi_dengan_pptx, fake_model
):
    row = sesi_dengan_pptx
    _seed_submitted(db, row)
    client = fake_model(
        FakeClient(
            [
                {
                    "skor_ideasi": 80,
                    "skor_isi": 60,
                    "skor_harapan": 50,
                    "skor_total": 100,  # bohong: wajib diabaikan
                    "ringkasan": "Jawaban Murid cukup detail pada aspek isi, "
                    "tapi tipis saat menjelaskan rencana lanjutan.",
                }
            ]
        )
    )

    tasks.assess_answers(row.id)

    with db() as session:
        penilaian = session.get(PenilaianAI, row.id)
        assert penilaian is not None
        assert (penilaian.skor_ideasi, penilaian.skor_isi, penilaian.skor_harapan) == (80, 60, 50)
        # round(mean(80, 60, 50)) = 63.33 — bukan 100 dari model.
        assert penilaian.skor_total == 63.33
        assert "skor" not in (penilaian.ringkasan or "").lower()
        assert penilaian.ringkasan.startswith("Jawaban Murid")

    # Idempoten: panggilan kedua tidak menyentuh model lagi.
    tasks.assess_answers(row.id)
    assert len(client.calls) == 1


def test_assess_answers_clamps_out_of_range_scores(db, sesi_dengan_pptx, fake_model):
    row = sesi_dengan_pptx
    _seed_submitted(db, row)
    fake_model(
        FakeClient(
            [
                {
                    "skor_ideasi": 150,
                    "skor_isi": -20,
                    "skor_harapan": 50,
                    "skor_total": 0,
                    "ringkasan": "Ringkasan penilaian untuk data yang keluar dari rentang.",
                }
            ]
        )
    )

    tasks.assess_answers(row.id)

    with db() as session:
        penilaian = session.get(PenilaianAI, row.id)
        assert penilaian.skor_ideasi == 100
        assert penilaian.skor_isi == 0
        assert penilaian.skor_total == 50.0


def test_assess_answers_skipped_when_not_submitted(db, sesi_dengan_pptx, fake_model, caplog):
    client = fake_model(FakeClient([]))

    with caplog.at_level(logging.INFO):
        tasks.assess_answers(sesi_dengan_pptx.id)

    assert client.calls == []
    with db() as session:
        assert session.get(PenilaianAI, sesi_dengan_pptx.id) is None


def test_assess_answers_skipped_when_answers_incomplete(db, sesi_dengan_pptx, fake_model):
    row = sesi_dengan_pptx
    _seed_submitted(db, row, with_answers=False)
    client = fake_model(FakeClient([]))

    tasks.assess_answers(row.id)

    assert client.calls == []
    with db() as session:
        assert session.get(PenilaianAI, row.id) is None


def test_assess_answers_swallows_exception(db, sesi_dengan_pptx, fake_model, caplog):
    row = sesi_dengan_pptx
    _seed_submitted(db, row)
    fake_model(FakeClient([RuntimeError("rate limit")]))

    with caplog.at_level(logging.ERROR):
        tasks.assess_answers(row.id)  # tidak melempar

    assert any("assess_answers gagal" in record.message for record in caplog.records)
    with db() as session:
        assert session.get(PenilaianAI, row.id) is None
        assert session.get(SessionRow, row.id).status == SessionStatus.SUBMITTED


def test_assess_document_prompt_lists_all_three_pairs():
    """Prompt penilaian memuat tiga pasang tipe + jawaban Murid."""
    pairs = [
        AnswerPair(QuestionType.IDEASI, "pertanyaan ideasi", "jawaban ideasi"),
        AnswerPair(QuestionType.ISI, "pertanyaan isi", "jawaban isi"),
        AnswerPair(QuestionType.HARAPAN, "pertanyaan harapan", "jawaban harapan"),
    ]
    client = FakeClient(
        [
            {
                "skor_ideasi": 70,
                "skor_isi": 70,
                "skor_harapan": 70,
                "skor_total": 70,
                "ringkasan": "Ringkasan netral untuk memeriksa konsistensi prompt penilaian.",
            }
        ]
    )

    result = asyncio.run(assess_document(client=client, brief=BRIEF, pairs=pairs))

    assert result.skor_total == 70.0
    assert result.declared_total == 70
    assert BRIEF in client.calls[0].prompt
    for tipe in QuestionType:
        assert tipe.value in client.calls[0].prompt
    assert "jawaban murid" in client.calls[0].prompt.lower()